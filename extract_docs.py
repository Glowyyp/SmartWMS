import os
import pandas as pd
import docx
from pptx import Presentation
import PyPDF2

base_dir = os.path.dirname(os.path.abspath(__file__))
evidencias_dir = os.path.join(base_dir, "Evidencias Proyecto")
out_dir = os.path.join(base_dir, "Evidencias_txt")
os.makedirs(out_dir, exist_ok=True)

for file in os.listdir(evidencias_dir):
    path = os.path.join(evidencias_dir, file)
    out_path = os.path.join(out_dir, file + ".txt")
    
    try:
        if file.endswith(".docx"):
            doc = docx.Document(path)
            text = "\n".join([para.text for para in doc.paragraphs])
            with open(out_path, "w", encoding="utf-8") as f:
                f.write(text)
        elif file.endswith(".xlsx"):
            xls = pd.ExcelFile(path)
            with open(out_path, "w", encoding="utf-8") as f:
                for sheet_name in xls.sheet_names:
                    df = pd.read_excel(xls, sheet_name=sheet_name)
                    f.write(f"\n--- SHEET: {sheet_name} ---\n")
                    f.write(df.to_csv(index=False))
        elif file.endswith(".pptx"):
            prs = Presentation(path)
            text = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text.append(shape.text)
            with open(out_path, "w", encoding="utf-8") as f:
                f.write("\n".join(text))
        elif file.endswith(".pdf"):
            reader = PyPDF2.PdfReader(path)
            text = []
            for page in reader.pages:
                text.append(page.extract_text())
            with open(out_path, "w", encoding="utf-8") as f:
                f.write("\n".join(text))
    except Exception as e:
        print(f"Error reading {file}: {e}")
        
print("Extraction complete.")
